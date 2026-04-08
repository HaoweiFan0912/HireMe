import base64
import mimetypes
import os
import subprocess
import tempfile
from io import BytesIO
from pathlib import Path
from typing import Any

from docx import Document
from pypdf import PdfReader

from app.core.paths import PDF_RENDER_SCRIPT, PROJECT_ROOT
from app.extraction.extraction_skills import (
    IMAGE_TEXT_EXTRACTION_SKILL,
    PDF_TEXT_PRIMARY_EXTRACTION_SKILL,
    SCANNED_PDF_EXTRACTION_SKILL,
)


def extract_pdf_text(file_bytes: bytes) -> str:
    reader = PdfReader(BytesIO(file_bytes))
    pages = []
    for page in reader.pages:
        page_text = page.extract_text() or ""
        if page_text.strip():
            pages.append(page_text.strip())
    return "\n\n".join(pages).strip()


def render_pdf_pages_to_images(file_bytes: bytes) -> list[bytes]:
    if not PDF_RENDER_SCRIPT.exists():
        raise ValueError("The PDF rendering script is missing, so scanned PDFs cannot be processed.")

    with tempfile.TemporaryDirectory(prefix="hireme_pdf_") as temp_dir:
        temp_path = Path(temp_dir)
        input_pdf = temp_path / "input.pdf"
        output_dir = temp_path / "rendered"
        input_pdf.write_bytes(file_bytes)

        completed = subprocess.run(
            ["swift", str(PDF_RENDER_SCRIPT), str(input_pdf), str(output_dir)],
            cwd=PROJECT_ROOT,
            capture_output=True,
            text=True,
        )
        if completed.returncode != 0:
            raise ValueError(
                "Scanned PDF rendering failed: "
                + (completed.stderr.strip() or completed.stdout.strip() or "unknown error")
            )

        images = [path.read_bytes() for path in sorted(output_dir.glob("*.png"))]
        if not images:
            raise ValueError("No page images were generated after rendering the scanned PDF.")
        return images


def extract_docx_text(file_bytes: bytes) -> str:
    document = Document(BytesIO(file_bytes))
    chunks: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            chunks.append(text)

    for table in document.tables:
        for row in table.rows:
            row_values = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if row_values:
                chunks.append(" | ".join(row_values))

    return "\n".join(chunks).strip()


def guess_mime_type(filename: str, content_type: str | None) -> str:
    return content_type or mimetypes.guess_type(filename)[0] or "application/octet-stream"


def decode_text(file_bytes: bytes) -> str:
    for encoding in ("utf-8", "utf-16", "latin-1"):
        try:
            return file_bytes.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ValueError("The text file encoding could not be detected.")


def build_document_content(
    filename: str,
    content_type: str | None,
    file_bytes: bytes,
) -> tuple[list[dict[str, Any]], str | None]:
    suffix = os.path.splitext(filename.lower())[1]
    mime_type = guess_mime_type(filename, content_type)

    if suffix == ".pdf" or mime_type == "application/pdf":
        text = extract_pdf_text(file_bytes)
        if text:
            content: list[dict[str, Any]] = [
                {
                    "type": "text",
                    "text": PDF_TEXT_PRIMARY_EXTRACTION_SKILL,
                },
                {
                    "type": "text",
                    "text": f"Document text:\n{text}",
                },
            ]
            try:
                rendered_images = render_pdf_pages_to_images(file_bytes)
            except Exception:
                rendered_images = []
            for image_bytes in rendered_images[:2]:
                data_url = f"data:image/png;base64,{base64.b64encode(image_bytes).decode('ascii')}"
                content.append(
                    {
                        "type": "image_url",
                        "image_url": {"url": data_url},
                    }
                )
            return content, text

        images = render_pdf_pages_to_images(file_bytes)
        content = [
            {
                "type": "text",
                "text": SCANNED_PDF_EXTRACTION_SKILL,
            }
        ]
        for image_bytes in images:
            data_url = f"data:image/png;base64,{base64.b64encode(image_bytes).decode('ascii')}"
            content.append(
                {
                    "type": "image_url",
                    "image_url": {"url": data_url},
                }
            )
        return content, None

    if suffix == ".docx" or mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        text = extract_docx_text(file_bytes)
        if not text:
            raise ValueError("No usable text was extracted from the DOCX file.")
        return ([{"type": "text", "text": f"Document text:\n{text}"}], text)

    if mime_type.startswith("text/") or suffix in {".txt", ".md", ".csv"}:
        text = decode_text(file_bytes)
        return ([{"type": "text", "text": f"Document text:\n{text}"}], text)

    if mime_type.startswith("image/") or suffix in {".png", ".jpg", ".jpeg", ".webp"}:
        data_url = f"data:{mime_type};base64,{base64.b64encode(file_bytes).decode('ascii')}"
        return (
            [
                {
                    "type": "text",
                    "text": IMAGE_TEXT_EXTRACTION_SKILL,
                },
                {
                    "type": "image_url",
                    "image_url": {"url": data_url},
                },
            ],
            None,
        )

    try:
        decoded = decode_text(file_bytes)
    except ValueError as exc:
        raise ValueError(
            "Unsupported file type. Please upload a PDF, DOCX, TXT, Markdown file, or image."
        ) from exc

    return ([{"type": "text", "text": f"Document text:\n{decoded}"}], decoded)
